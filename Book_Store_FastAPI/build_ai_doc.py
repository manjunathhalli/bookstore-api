"""Generate a Google-Docs-ready .docx documenting the AI features we added to
the FastAPI Book Store — the native port of the Laravel AI Integration Guide.

Run:  python build_ai_doc.py
Output: Book_Store_FastAPI_AI_Features.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK = RGBColor(0x1E, 0x29, 0x3B)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
GREY = RGBColor(0x55, 0x5F, 0x70)
CODE_BG = "F2F3F7"


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, CODE_BG)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Consolas")
    rfonts.set(qn("w:hAnsi"), "Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def body(doc, text, *, italic=False, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(8)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text).font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text).font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)


def make_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].paragraphs[0].add_run(h).bold = True
        for r in hdr[i].paragraphs[0].runs:
            r.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].add_run(val).font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def h1(text):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = DARK


def h2(text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = INDIGO


# ---- Title ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Book Store (FastAPI)")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = DARK
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("AI Features with the Claude API — All 10 Concepts")
rs.font.size = Pt(14)
rs.font.color.rgb = INDIGO
intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
ri = intro.add_run("A native FastAPI port of the Laravel AI Integration Guide — concepts, setup, and the actual code.")
ri.italic = True
ri.font.size = Pt(10.5)
ri.font.color.rgb = GREY
doc.add_paragraph()


# ===== 1 =====
h1("1. Overview")
body(doc, "This document explains the Artificial Intelligence features added to the FastAPI "
          "Book Store. It is a faithful port of the Laravel 'AI Integration Guide': the same "
          "10 AI concepts, powered by the same Anthropic Claude API, but implemented the "
          "FastAPI way — feature-first packages, SQLAlchemy models, Pydantic schemas, and a "
          "shared service class.")
body(doc, "Like the rest of the project, every AI feature is exposed through BOTH doors: a "
          "JWT REST API (under /api/ai) and the session-based web UI (under /ai). All AI calls "
          "happen server-side; the API key never reaches the browser. Every feature degrades "
          "gracefully when no key is configured.")
make_table(
    doc,
    ["#", "AI Concept", "Kind of AI"],
    [
        ["1", "Sentiment Analysis", "Text classification"],
        ["2", "Text Generation (descriptions)", "Content generation"],
        ["3", "Summarization", "Condensing many docs into one"],
        ["4", "Natural-Language Search", "Intent parsing to JSON"],
        ["5", "AI Chatbot (RAG)", "Retrieval-Augmented Generation"],
        ["6", "Personalized Recommendations", "Reasoning over user history"],
        ["7", "Content Moderation", "Safety classification"],
        ["8", "Auto-Tagging / Classification", "Multi-label classification"],
        ["9", "Demand Forecasting", "Reasoning over statistics"],
        ["10", "Semantic Search", "Embeddings + cosine similarity"],
    ],
)


# ===== 2 =====
h1("2. The 10 AI Concepts Explained")
body(doc, "Before the code, here is what each concept means and how we use it in the store.")

concepts = [
    ("Sentiment Analysis",
     "The model reads a review and labels its emotional tone (Positive / Neutral / Negative). "
     "We store the label in a new feedbacks.sentiment column so admins can tell happy from "
     "unhappy customers at a glance."),
    ("Text Generation",
     "The model writes new content — a marketing description — from a short prompt (title + "
     "author). Used by admins when adding a book."),
    ("Summarization",
     "The model condenses many input documents (all reviews of a book) into a short, balanced "
     "2-3 sentence summary."),
    ("Natural-Language Search (Intent Parsing)",
     "The model turns a free-text query like 'cheap thrillers under 500' into a structured JSON "
     "filter, which we then apply with ordinary SQLAlchemy queries — never raw SQL."),
    ("Retrieval-Augmented Generation (RAG)",
     "We first retrieve relevant books from OUR database, then pass them to the model as context, "
     "so the chatbot answers using the real catalogue instead of making things up."),
    ("Personalized Recommendation",
     "The model reasons over a user's orders and wishlist to suggest up to 4 books they'll enjoy, "
     "each with a one-sentence reason."),
    ("Content Moderation",
     "The model checks user text for hate, harassment, sexual content or spam and returns SAFE / "
     "UNSAFE, letting us reject bad reviews before they are stored."),
    ("Auto-Tagging / Classification",
     "The model assigns 1-4 lowercase genre tags to a book from its title and description, stored "
     "as a JSON column for browsing and filtering."),
    ("Demand Forecasting",
     "We aggregate real order counts per book in SQL, then let the model reason over the numbers "
     "and current stock to produce a ranked, restock-flagged list. Hard numbers from the DB; "
     "interpretation from the model."),
    ("Semantic Search (Embeddings)",
     "Text is converted into numeric vectors; we compare a query vector to each book's stored "
     "vector by cosine similarity, so a search matches by MEANING, not just keywords."),
]
for name, desc in concepts:
    h2(name)
    body(doc, desc)


# ===== 3 =====
h1("3. Anthropic Account & Subscription Setup")
body(doc, "All AI features call the Anthropic Claude API. You need an API account with credits — "
          "this is separate from a Claude.ai chat subscription.")
h2("Step 3.1 — Create a Console account")
bullet(doc, "Go to https://console.anthropic.com and sign up. Verify your email.")
h2("Step 3.2 — Add billing / credits")
bullet(doc, "The API is pay-as-you-go. Open Billing and purchase a small amount of prepaid "
            "credits (e.g. $5) to start.")
bullet(doc, "A Claude.ai Pro chat subscription is NOT the same as API credits. For code you "
            "need API credits in console.anthropic.com.")
h2("Step 3.3 — Create an API key")
bullet(doc, "Settings -> API Keys -> Create Key. Copy it immediately (shown only once).")
bullet(doc, "Never commit the key to git. It lives only in .env, which is gitignored.")
h2("Step 3.4 — Embeddings for Semantic Search (local, no API key)")
bullet(doc, "Semantic Search (Feature 10) needs embeddings, but Anthropic does NOT offer an "
            "embeddings API. Rather than depend on a third-party service (the original guide used "
            "Voyage AI), this port runs a local sentence-transformers model — no key, no external "
            "call. It ships in requirements.txt (sentence-transformers); the ~80 MB model is "
            "downloaded once from Hugging Face on first use and cached, then runs fully offline. "
            "The other 9 features are unaffected if it is not installed.")


# ===== 4 =====
h1("4. FastAPI Project Setup")
h2("Step 4.1 — Add configuration to .env")
body(doc, "Copy from .env.example and fill in your keys. Model IDs mirror the guide and are "
          "configurable:")
code_block(doc, [
    "# ---- AI features (Anthropic Claude API) ----",
    "ANTHROPIC_API_KEY=sk-ant-xxxxxxxxxxxxxxxxxxxx",
    "ANTHROPIC_MODEL=claude-sonnet-4-6",
    "ANTHROPIC_MODEL_FAST=claude-haiku-4-5-20251001",
    "",
    "# ---- Embeddings (Semantic Search) ----",
    "# No key needed: a local sentence-transformers model is used.",
    "# Just install requirements.txt (includes sentence-transformers).",
])
h2("Step 4.2 — Register the settings")
body(doc, "In app/core/config.py the Settings class reads them (with safe defaults). Optional keys "
          "mean the app still boots with AI disabled:")
code_block(doc, [
    "class Settings(BaseSettings):",
    "    ...",
    "    anthropic_api_key: str = ''",
    "    anthropic_model: str = 'claude-sonnet-4-6'",
    "    anthropic_model_fast: str = 'claude-haiku-4-5-20251001'",
    "    anthropic_base_url: str = 'https://api.anthropic.com/v1'",
    "    anthropic_version: str = '2023-06-01'",
    "    # (Voyage settings retired — embeddings now run locally, no key.)",
])
h2("Step 4.3 — HTTP client choice")
body(doc, "FastAPI already ships httpx. We call the Claude REST API with it directly — no extra "
          "SDK — mirroring the guide's use of Laravel's Http client. Embeddings run in-process via "
          "sentence-transformers (no HTTP). Route handlers are plain 'def', so blocking calls run "
          "in Starlette's threadpool and never stall the event loop.")


# ===== 5 =====
h1("5. Core: The ClaudeService Class")
body(doc, "One reusable client every feature calls — the FastAPI equivalent of Laravel's "
          "ClaudeService. It lives in app/core/ai.py and offers ask() (text), ask_json() "
          "(parsed JSON), and embed() (a local sentence-transformers vector), plus capability "
          "flags the UI uses to show/hide AI buttons.")
code_block(doc, [
    "# app/core/ai.py",
    "class ClaudeService:",
    "    def __init__(self):",
    "        self._key       = settings.anthropic_api_key",
    "        self._base_url  = settings.anthropic_base_url.rstrip('/')",
    "        self._version   = settings.anthropic_version",
    "        self._model     = settings.anthropic_model",
    "        self._model_fast= settings.anthropic_model_fast",
    "",
    "    @property",
    "    def enabled(self) -> bool:            # Claude key present?",
    "        return bool(self._key)",
    "    @property",
    "    def embeddings_enabled(self) -> bool: # local model loadable?",
    "        return self._get_embed_model() is not None",
    "",
    "    def ask(self, prompt, system=None, model=None, max_tokens=1024) -> str:",
    "        if not self._key:",
    "            raise AIError('ANTHROPIC_API_KEY is not configured.')",
    "        payload = {'model': model or self._model,",
    "                   'max_tokens': max_tokens,",
    "                   'messages': [{'role': 'user', 'content': prompt}]}",
    "        if system: payload['system'] = system",
    "        resp = httpx.post(f'{self._base_url}/messages', headers={",
    "                   'x-api-key': self._key,",
    "                   'anthropic-version': self._version,",
    "                   'content-type': 'application/json'},",
    "                   json=payload, timeout=60)",
    "        if resp.status_code >= 400:",
    "            raise AIError(f'AI request failed: {resp.status_code}')",
    "        blocks = resp.json().get('content') or []",
    "        return blocks[0].get('text', '') if blocks else ''",
    "",
    "    def ask_json(self, prompt, system=None, model=None, max_tokens=1024):",
    "        raw = self.ask(prompt, system, model, max_tokens).strip()",
    "        raw = re.sub(r'^```(?:json)?|```$', '', raw, flags=re.M).strip()",
    "        try:    return json.loads(raw)",
    "        except (json.JSONDecodeError, ValueError): return {}",
    "",
    "    _EMBED_MODEL_NAME = 'all-MiniLM-L6-v2'",
    "    _embed_model = None                     # loaded once, shared",
    "",
    "    @classmethod",
    "    def _get_embed_model(cls):              # None if package missing",
    "        if cls._embed_model is None:",
    "            try:",
    "                from sentence_transformers import SentenceTransformer",
    "            except ImportError:",
    "                return None",
    "            cls._embed_model = SentenceTransformer(cls._EMBED_MODEL_NAME)",
    "        return cls._embed_model",
    "",
    "    def embed(self, text) -> list[float]:   # local sentence-transformers",
    "        model = self._get_embed_model()",
    "        if model is None:",
    "            raise AIError('sentence-transformers is not installed.')",
    "        return model.encode(text, normalize_embeddings=True).tolist()",
    "",
    "    # Retired: Voyage AI embeddings (Anthropic has no embeddings API);",
    "    # the original httpx call is kept, commented out, in app/core/ai.py.",
    "",
    "claude = ClaudeService()          # one shared instance",
    "def get_claude(): return claude   # FastAPI dependency / accessor",
])
body(doc, "Any endpoint gets it with claude: ClaudeService = Depends(get_claude). The reasoning "
          "for each feature lives in app/ai/service.py; the HTTP surface in app/ai/web.py and "
          "app/ai/api.py.", italic=True)


# ===== 6 =====
h1("6. Database: The Three AI Columns")
body(doc, "Three AI features persist data, so we added three columns to the existing models:")
make_table(
    doc,
    ["Column", "Table", "Type", "Feature"],
    [
        ["sentiment", "feedbacks", "VARCHAR(20)", "1 — Sentiment analysis"],
        ["tags", "books", "JSON", "8 — Auto-tagging"],
        ["embedding", "books", "JSON", "10 — Semantic search"],
    ],
)
body(doc, "Because Base.metadata.create_all only CREATEs missing tables (never ALTERs an "
          "existing one), a helper adds these columns idempotently on startup when running "
          "against the pre-existing Laravel schema:")
code_block(doc, [
    "# app/core/database.py",
    "_AI_COLUMNS = {",
    "    'feedbacks': [('sentiment', 'VARCHAR(20) NULL')],",
    "    'books':     [('tags', 'JSON NULL'), ('embedding', 'JSON NULL')],",
    "}",
    "def ensure_ai_columns():",
    "    insp = inspect(engine)",
    "    with engine.begin() as conn:",
    "        for table, cols in _AI_COLUMNS.items():",
    "            if table not in insp.get_table_names(): continue",
    "            present = {c['name'] for c in insp.get_columns(table)}",
    "            for name, ddl in cols:",
    "                if name in present: continue",
    "                try: conn.execute(text(f'ALTER TABLE {table} ADD COLUMN {name} {ddl}'))",
    "                except Exception: pass   # never block startup",
    "",
    "# called from main.py on_startup(), right after create_all(...)",
])


# ===== Feature sections =====
h1("7. Feature 1 — Sentiment Analysis on Reviews")
body(doc, "When a review is saved we ask the fast model for a one-word label and store it. The "
          "helper lives in app/ai/service.py:")
code_block(doc, [
    "def analyze_sentiment(claude, feedback_text):",
    "    system = ('You are a sentiment classifier. Reply with exactly one word: '",
    "              'Positive, Neutral, or Negative.')",
    "    try:",
    "        label = claude.ask(f'Review: \"{feedback_text}\"', system,",
    "                           claude.fast_model, max_tokens=5).strip()",
    "    except AIError:",
    "        return None",
    "    return label if label in {'Positive','Neutral','Negative'} else 'Neutral'",
])
body(doc, "It is wired into the feedback store (both web.py and api.py) alongside moderation "
          "(Feature 7) — see Section 13.")

h1("8. Feature 2 — AI Book Description Generator (Admin)")
body(doc, "An admin clicks '✨ Generate with AI' on the add-book form; JavaScript posts the name "
          "and author and drops the returned text into the description box.")
code_block(doc, [
    "# service.py",
    "def generate_description(claude, name, author):",
    "    system = ('You are a bookstore copywriter. Write one engaging product '",
    "              'description of 40-60 words. No headings.')",
    "    return claude.ask(f'Write a description for \"{name}\" by {author}.',",
    "                      system, max_tokens=300).strip()",
    "",
    "# web.py  (admin-only, called via fetch)",
    "@router.post('/books/generate-description')",
    "def generate_description(payload: DescriptionRequest,",
    "        user=Depends(admin_role), claude=Depends(get_claude)):",
    "    if not claude.enabled:",
    "        return JSONResponse({'error': 'AI is not configured.'}, status_code=503)",
    "    text = service.generate_description(claude, payload.name, payload.author)",
    "    return JSONResponse({'description': text})",
])

h1("9. Feature 3 — Review Summarization")
body(doc, "A '✨ AI Review Summary' button on the feedback page condenses up to 50 reviews of a "
          "book into a balanced summary.")
code_block(doc, [
    "def summarize_reviews(claude, db, book_id):",
    "    reviews = db.scalars(select(Feedback.feedback)",
    "        .where(Feedback.book_id == book_id,",
    "               Feedback.feedback.isnot(None)).limit(50)).all()",
    "    if not reviews:",
    "        return 'No reviews yet.'",
    "    joined = '\\n'.join(f'{i+1}. {r}' for i, r in enumerate(reviews))",
    "    system = ('Summarize customer reviews in 2-3 sentences. Mention common '",
    "              'praises and complaints. Be balanced.')",
    "    return claude.ask(f'Reviews:\\n{joined}', system, max_tokens=250).strip()",
])
body(doc, "Cap at 50 reviews to control cost; cache popular summaries if traffic grows.", italic=True)

h1("10. Feature 4 — Natural-Language Search")
body(doc, "The model returns a STRUCTURED FILTER (not SQL); we apply it with SQLAlchemy bindings, "
          "so there is no injection risk.")
code_block(doc, [
    "def parse_search_filters(claude, query):",
    "    system = ('You convert a book-store search query into JSON filters.\\n'",
    "              'Return ONLY JSON with optional keys:\\n'",
    "              '  \"keywords\",\"author\",\"max_price\",\"min_price\",\\n'",
    "              '  \"sort\" (\"price_asc\"|\"price_desc\"|\"newest\"). Omit unused keys.')",
    "    f = claude.ask_json(query, system, claude.fast_model)",
    "    return f if isinstance(f, dict) else {}",
    "",
    "def smart_search(claude, db, query):",
    "    f = parse_search_filters(claude, query)",
    "    stmt = select(Book)",
    "    if f.get('keywords'):",
    "        like = f\"%{f['keywords']}%\"",
    "        stmt = stmt.where((Book.name.like(like)) | (Book.description.like(like)))",
    "    if f.get('author'):",
    "        stmt = stmt.where(Book.author.like(f\"%{f['author']}%\"))",
    "    books = list(db.scalars(stmt).all())",
    "    # price is a string column -> compare/sort numerically in Python",
    "    ... apply max_price / min_price / sort ...",
    "    return books, f",
])

h1("11. Feature 5 — AI Chatbot (RAG)")
body(doc, "We retrieve up to 15 matching books, build a small catalogue string, and constrain the "
          "model to answer only from it — Retrieval-Augmented Generation.")
code_block(doc, [
    "def chat_reply(claude, db, message):",
    "    like = f'%{message}%'",
    "    books = db.scalars(select(Book).where(",
    "        (Book.name.like(like)) | (Book.author.like(like)) |",
    "        (Book.description.like(like))).limit(15)).all()",
    "    if not books:",
    "        books = db.scalars(select(Book).order_by(Book.name).limit(15)).all()",
    "    catalog = '\\n'.join(",
    "        f'- {b.name} by {b.author} (Rs.{b.price}): {(b.description or \"\")[:120]}'",
    "        for b in books)",
    "    system = ('You are \"BookBot\", a friendly bookstore assistant.\\n'",
    "              'Answer ONLY using the catalog below. Recommend specific titles '",
    "              f'with price.\\nCATALOG:\\n{catalog}')",
    "    return claude.ask(message, system, max_tokens=600).strip()",
])
body(doc, "The web page (/ai/chat) posts messages via fetch and appends the reply as a chat bubble.")

h1("12. Feature 6 — Personalized Recommendations")
body(doc, "We gather the user's ordered and wishlisted titles, hand the catalogue to the model, "
          "and ask for a JSON array of picks with reasons — validated against real IDs before use.")
code_block(doc, [
    "def recommend(claude, db, user_id):",
    "    ordered = db.scalars(select(Book.name).join(Order, Order.book_id==Book.id)",
    "        .where(Order.user_id==user_id).distinct()).all()",
    "    wished  = db.scalars(select(Book.name).join(WishList, WishList.book_id==Book.id)",
    "        .where(WishList.user_id==user_id).distinct()).all()",
    "    catalog = db.scalars(select(Book)).all()",
    "    by_id = {b.id: b for b in catalog}",
    "    listing = '\\n'.join(f'{b.id}: {b.name} by {b.author}' for b in catalog)",
    "    system = ('You are a book recommender. From the CATALOG pick up to 4 books '",
    "              'the user will enjoy. Return ONLY a JSON array: '",
    "              '[{\"id\":<id>,\"reason\":\"<sentence>\"}]. '",
    "              'Do not recommend books already ordered.')",
    "    picks = claude.ask_json(..., system)",
    "    # keep only picks whose id exists in our catalog",
    "    return [{'book': by_id[p['id']], 'reason': p.get('reason','')} ",
    "            for p in picks if by_id.get(p.get('id'))]",
])

h1("13. Feature 7 — Content Moderation")
body(doc, "Before saving any user review we check it. Fail-OPEN — an API outage never blocks a "
          "user over a non-critical review; skipped entirely if no key is set.")
code_block(doc, [
    "def moderate(claude, text_value) -> bool:",
    "    if not claude.enabled:",
    "        return True",
    "    system = ('You are a content moderator for a bookstore review section. '",
    "              'Decide if the text contains hate, harassment, sexual content, '",
    "              'or spam. Reply with ONLY the word SAFE or UNSAFE.')",
    "    try:",
    "        verdict = claude.ask(f'Text: \"{text_value}\"', system,",
    "                             claude.fast_model, max_tokens=5).strip()",
    "    except AIError:",
    "        return True   # fail-open",
    "    return verdict.upper() == 'SAFE'",
])
body(doc, "Moderation + sentiment are wired straight into the existing feedback store:")
code_block(doc, [
    "# app/feedback/web.py  (inside store())",
    "claude = get_claude()",
    "if not ai.moderate(claude, feedback):",
    "    flash_errors(request, ['Your review was flagged as inappropriate.'],",
    "                 {'feedback': feedback})",
    "    return RedirectResponse('/feedback', status_code=303)",
    "sentiment = ai.analyze_sentiment(claude, feedback) if claude.enabled else None",
    "db.add(Feedback(user_id=user.id, book_id=book_id, feedback=feedback,",
    "                rating=rating, sentiment=sentiment))",
])

h1("14. Feature 8 — Auto-Tagging / Categorization")
body(doc, "When a book is created or updated we ask the fast model for 1-4 genre tags and store "
          "them in the JSON column.")
code_block(doc, [
    "def auto_tag(claude, name, description) -> list[str]:",
    "    system = ('You tag books by genre. From the title and description, return '",
    "              'ONLY a JSON array of 1-4 lowercase genre tags, e.g. '",
    "              '[\"fiction\",\"thriller\",\"mystery\"].')",
    "    tags = claude.ask_json(f'Title: {name}\\nDescription: {description}',",
    "                           system, claude.fast_model)",
    "    return [t for t in tags if isinstance(t, str)][:4]",
    "",
    "# app/books/web.py — best-effort enrichment on create/update",
    "def _ai_enrich(book):",
    "    claude = get_claude()",
    "    if claude.enabled:",
    "        book.tags = ai.auto_tag(claude, book.name, book.description)",
    "    if claude.embeddings_enabled:",
    "        try: book.embedding = ai.embed_book_text(claude, book.name,",
    "                                   book.author, book.description)",
    "        except AIError: pass",
])

h1("15. Feature 9 — Demand Forecasting")
body(doc, "Real statistics + AI reasoning. We aggregate order counts per book in SQL, then let the "
          "model rank titles and flag restocks. The hard numbers keep the result grounded.")
code_block(doc, [
    "def forecast(claude, db):",
    "    rows = db.execute(select(Book.name,",
    "            func.count(Order.id).label('orders_count'),",
    "            Book.quantity.label('stock'))",
    "        .join(Order, Order.book_id == Book.id)",
    "        .group_by(Book.id, Book.name, Book.quantity)",
    "        .order_by(func.count(Order.id).desc()).limit(30)).all()",
    "    if not rows: return []",
    "    data = '\\n'.join(f'{r.name}: {r.orders_count} orders, {r.stock} in stock'",
    "                     for r in rows)",
    "    system = ('You are a demand-forecasting analyst... return ONLY a JSON array '",
    "              'of up to 8 objects: [{\"book\":...,\"trend\":\"rising|steady|falling\",'",
    "              '\"restock\":true|false,\"reason\":...}]. Flag restock when demand is '",
    "              'high but stock is low.')",
    "    return claude.ask_json(f'DATA:\\n{data}', system)",
])

h1("16. Feature 10 — Semantic Search (Embeddings)")
body(doc, "Search by MEANING. Each book stores a local sentence-transformers embedding (384-dim); a "
          "query is embedded the same way and ranked by cosine similarity — computed in Python, "
          "which is fine for a few thousand books.")
code_block(doc, [
    "def cosine(a, b):",
    "    if not a or not b: return 0.0",
    "    dot = sum(x*y for x, y in zip(a, b))",
    "    na  = math.sqrt(sum(x*x for x in a))",
    "    nb  = math.sqrt(sum(y*y for y in b))",
    "    return dot/(na*nb) if na and nb else 0.0",
    "",
    "def semantic_search(claude, db, query, limit=12):",
    "    q = claude.embed(query)",
    "    if not q: return []",
    "    scored = []",
    "    for book in db.scalars(select(Book).where(Book.embedding.isnot(None))).all():",
    "        book.score = cosine(q, book.embedding or [])",
    "        scored.append((book.score, book))",
    "    scored.sort(key=lambda p: p[0], reverse=True)",
    "    return [b for _, b in scored[:limit]]",
])
body(doc, "After installing sentence-transformers, run 'Reindex embeddings' on the Semantic Search "
          "page (or POST /api/ai/reindex-embeddings) to vectorise books created before the model "
          "was available. For much larger catalogues, move vectors to a dedicated store "
          "(pgvector, Qdrant).",
     italic=True)


# ===== Endpoints map =====
h1("17. Endpoint Map (Web + JWT API)")
make_table(
    doc,
    ["Feature", "Web (session)", "API (JWT)"],
    [
        ["Hub", "GET /ai", "—"],
        ["Chatbot (RAG)", "GET/POST /ai/chat", "POST /api/ai/chat"],
        ["NL search", "GET /ai/search", "POST /api/ai/smart-search"],
        ["Semantic search", "GET /ai/semantic", "POST /api/ai/semantic-search"],
        ["Reindex embeddings", "POST /ai/reindex-embeddings", "POST /api/ai/reindex-embeddings"],
        ["Recommendations", "GET /ai/recommendations", "GET /api/ai/recommendations"],
        ["Forecast (admin)", "GET /ai/forecast", "GET /api/ai/forecast"],
        ["Description (admin)", "POST /ai/books/generate-description", "POST /api/ai/generate-description"],
        ["Review summary", "POST /ai/reviews/summary", "GET /api/ai/reviews/{id}/summary"],
        ["Auto-tag (admin)", "(on book save)", "POST /api/ai/auto-tag"],
        ["Sentiment + moderation", "(on feedback save)", "(on feedback save)"],
    ],
)


# ===== Security =====
h1("18. Security & Best Practices")
bullet(doc, "Key stays server-side. Every Claude call happens in the backend; the browser only "
            "talks to our own /ai routes.")
bullet(doc, "Keys out of git — they live only in .env (gitignored); .env.example is the template.")
bullet(doc, "Wrap AI calls so an outage degrades gracefully (AIError caught; features return "
            "safe fallbacks).")
bullet(doc, "Validate and constrain model output — accept only known labels/keys; never run "
            "model output as code or SQL. Filters flow through SQLAlchemy bindings.")
bullet(doc, "Use the cheapest capable model — Haiku (fast) for classification / moderation / "
            "parsing / tagging; Sonnet for generation and reasoning.")
bullet(doc, "Cache expensive results (summaries, recommendations, embeddings) to cut cost.")
bullet(doc, "Consider rate-limiting the AI routes to prevent runaway cost.")


# ===== Testing =====
h1("19. Testing & Verification")
body(doc, "Quick smoke test of the client from a Python shell:")
code_block(doc, [
    "python -c \"from app.core.ai import claude; print(claude.ask('Say hello in 3 words.'))\"",
])
body(doc, "During development the service layer was exercised end-to-end against an in-memory "
          "SQLite database with a mocked ClaudeService — all 10 helpers (sentiment, moderation, "
          "description, summary, smart-search filtering, RAG chat, recommendations, auto-tag, "
          "forecast, cosine-ranked semantic search) returned correct results, and the column "
          "migration was confirmed idempotent.")
body(doc, "Suggested rollout order: keys + ClaudeService -> Sentiment + Moderation (reuse the "
          "feedback flow) -> Description, Summary, Auto-Tagging (admin) -> NL Search -> Chatbot "
          "-> Recommendations -> Forecasting -> Semantic Search (embeddings, last).", italic=True)


# ===== Summary =====
h1("One-Paragraph Summary")
body(doc, "One shared ClaudeService (app/core/ai.py) wraps the Claude API for text features and a "
          "local sentence-transformers model for embeddings (Anthropic has no embeddings API). The "
          "reasoning for all 10 features lives in app/ai/service.py, exposed through both a web "
          "router (/ai) and a JWT API router (/api/ai). Three concepts persist data via three "
          "new columns (feedbacks.sentiment, books.tags, books.embedding) added idempotently on "
          "startup. Sentiment and moderation hook into the feedback flow; auto-tagging and "
          "embedding hook into the book flow; the rest are their own pages/endpoints. Every call "
          "is server-side, output is validated, and missing keys disable features gracefully — "
          "the same architecture as the Laravel guide, done the FastAPI way.", italic=True)

out = "Book_Store_FastAPI_AI_Features.docx"
doc.save(out)
print("Saved:", out)

"""Generate a Google-Docs-ready .docx of the *Backend-Only API* build guide.

Run:  python build_api_doc.py
Output: Book_Store_FastAPI_API_From_Scratch.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK = RGBColor(0x1E, 0x29, 0x3B)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
GREY = RGBColor(0x55, 0x5F, 0x70)
CODE_BG = "F2F3F7"


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, CODE_BG)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Consolas")
    rfonts.set(qn("w:hAnsi"), "Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def body(doc, text, *, italic=False, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(8)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text).font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text).font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)


def make_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].paragraphs[0].add_run(h).bold = True
        for r in hdr[i].paragraphs[0].runs:
            r.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].add_run(val).font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def h1(text):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = DARK


def h2(text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = INDIGO


# ---- Title ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Book Store API (FastAPI)")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = DARK
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("Backend-Only — Build From Scratch, Step by Step")
rs.font.size = Pt(14)
rs.font.color.rgb = INDIGO
intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
ri = intro.add_run("Pure JSON REST API protected by a JWT bearer token — no HTML, "
                   "templates, sessions, or file uploads.")
ri.italic = True
ri.font.size = Pt(10.5)
ri.font.color.rgb = GREY
doc.add_paragraph()

body(doc, "This guide builds ONLY the JSON REST API backend. Layout is feature-first: shared "
          "infrastructure in app/core, shared JWT guards in app/auth, and one package per "
          "domain holding its models.py, schemas.py, api.py. Type the code yourself as you go.")

# ===== 0 =====
h1("0. What 'Backend-Only' Changes")
body(doc, "Compared to a full-stack build, the API-only version drops these:")
make_table(doc, ["Dropped", "Why"], [
    ["web.py routers", "No server-rendered HTML pages."],
    ["templates/, Jinja2", "Nothing to render."],
    ["SessionMiddleware, cookies", "Auth is stateless JWT, not a login cookie."],
    ["dashboard/, users/, password/ web packages", "Those were web-only screens."],
    ["python-multipart, file uploads", "Book covers are a string URL, not a file."],
    ["core/templating.py, RedirectException", "Web-only concerns."],
])
body(doc, "What stays: config, database, security (JWT), the 7 tables, schemas, JWT guards, "
          "and the seven API routers.")

# ===== 1 =====
h1("1. Prerequisites")
make_table(doc, ["Need", "Why", "Check"], [
    ["Python 3.11+", "Modern typing (str | None, Mapped[...])", "python --version"],
    ["MySQL running", "Stores the data (XAMPP is fine)", "XAMPP -> Start MySQL"],
    ["Postman or curl", "To call the endpoints", "-"],
])

# ===== 2 =====
h1("2. Create the Project Skeleton")
code_block(doc, [
    "mkdir Book_Store_API",
    "cd Book_Store_API",
    "",
    "python -m venv .venv",
    ".\\.venv\\Scripts\\Activate.ps1",
    "",
    "mkdir app",
    "mkdir app\\core",
    "mkdir app\\auth app\\books app\\cart app\\wishlist app\\address app\\orders app\\feedback",
])
body(doc, "Add an empty __init__.py to every package:")
code_block(doc, [
    "ni app\\__init__.py",
    "ni app\\core\\__init__.py",
    "foreach ($f in 'auth','books','cart','wishlist','address','orders','feedback')",
    "  { ni app\\$f\\__init__.py }",
])
body(doc, "Final structure:")
code_block(doc, [
    "Book_Store_API/",
    "|-- run.py",
    "|-- requirements.txt",
    "|-- .env  /  .env.example",
    "|-- app/",
    "    |-- main.py",
    "    |-- core/   config.py database.py security.py",
    "    |-- auth/   models.py schemas.py dependencies.py api.py",
    "    |-- books/  models.py schemas.py api.py",
    "    |-- cart/ wishlist/ address/ orders/ feedback/   (same shape)",
])

# ===== 3 =====
h1("3. Install the Libraries")
body(doc, "requirements.txt (no jinja2 / multipart / itsdangerous — those were web-only):")
code_block(doc, [
    "fastapi", "uvicorn[standard]", "sqlalchemy", "pymysql",
    "pydantic[email]", "pydantic-settings", "passlib[bcrypt]", "pyjwt",
])
code_block(doc, ["pip install -r requirements.txt"])
make_table(doc, ["Library", "Role"], [
    ["fastapi", "Routes, validation, dependency injection, auto Swagger docs."],
    ["uvicorn", "ASGI server that runs the app."],
    ["sqlalchemy", "ORM — Python classes instead of raw SQL."],
    ["pymysql", "Driver SQLAlchemy uses to reach MySQL."],
    ["pydantic / -settings", "Validates request bodies; loads settings from .env."],
    ["passlib[bcrypt]", "Hashes passwords."],
    ["pyjwt", "Creates/reads the JWT bearer token."],
])

# ===== 4 =====
h1("4. Configuration — app/core/config.py")
code_block(doc, [
    "from functools import lru_cache",
    "from pydantic_settings import BaseSettings, SettingsConfigDict",
    "",
    "class Settings(BaseSettings):",
    "    app_name: str = 'Book Store API'",
    "    secret_key: str = 'insecure-dev-secret-change-me'   # signs the JWT",
    "    access_token_expire_minutes: int = 1440",
    "    algorithm: str = 'HS256'",
    "    db_host: str = '127.0.0.1'; db_port: int = 3306",
    "    db_database: str = 'book_store_product'",
    "    db_username: str = 'root'; db_password: str = ''",
    "    create_tables: bool = True",
    "    model_config = SettingsConfigDict(env_file='.env', extra='ignore')",
    "",
    "    @property",
    "    def database_url(self) -> str:",
    "        return (f'mysql+pymysql://{self.db_username}:{self.db_password}'",
    "                f'@{self.db_host}:{self.db_port}/{self.db_database}?charset=utf8mb4')",
    "",
    "@lru_cache",
    "def get_settings(): return Settings()",
    "settings = get_settings()",
])
body(doc, ".env (and a committable .env.example with the same keys):")
code_block(doc, [
    "APP_NAME=Book Store API",
    "SECRET_KEY=change-me-to-a-long-random-string",
    "DB_HOST=127.0.0.1", "DB_PORT=3306",
    "DB_DATABASE=book_store_product",
    "DB_USERNAME=root", "DB_PASSWORD=", "CREATE_TABLES=true",
])

# ===== 5 =====
h1("5. Database Engine — app/core/database.py")
code_block(doc, [
    "from collections.abc import Generator",
    "from datetime import datetime",
    "from sqlalchemy import DateTime, create_engine, func",
    "from sqlalchemy.orm import (DeclarativeBase, Mapped, Session,",
    "                            mapped_column, sessionmaker)",
    "from .config import settings",
    "",
    "engine = create_engine(settings.database_url, pool_pre_ping=True, future=True)",
    "SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False, future=True)",
    "",
    "class Base(DeclarativeBase): pass",
    "",
    "class TimestampMixin:",
    "    created_at: Mapped[datetime | None] = mapped_column(",
    "        DateTime, server_default=func.now())",
    "    updated_at: Mapped[datetime | None] = mapped_column(",
    "        DateTime, server_default=func.now(), onupdate=func.now())",
    "",
    "def get_db() -> Generator[Session, None, None]:",
    "    db = SessionLocal()",
    "    try:    yield db",
    "    finally: db.close()",
])
body(doc, "Every endpoint that touches the DB takes db: Session = Depends(get_db). FastAPI "
          "opens the session, hands it over, and the finally always closes it.", italic=True)

# ===== 6 =====
h1("6. Passwords + JWT — app/core/security.py")
code_block(doc, [
    "from datetime import datetime, timedelta, timezone",
    "import jwt",
    "from passlib.context import CryptContext",
    "from .config import settings",
    "",
    "pwd_context = CryptContext(schemes=['bcrypt'], deprecated='auto')",
    "",
    "def hash_password(plain): return pwd_context.hash(plain)",
    "def verify_password(plain, hashed):",
    "    try:    return pwd_context.verify(plain, hashed)",
    "    except ValueError: return False",
    "",
    "def create_access_token(subject, expires_minutes=None):",
    "    expire = datetime.now(timezone.utc) + timedelta(",
    "        minutes=expires_minutes or settings.access_token_expire_minutes)",
    "    return jwt.encode({'sub': str(subject), 'exp': expire},",
    "                      settings.secret_key, algorithm=settings.algorithm)",
    "",
    "def decode_access_token(token):",
    "    try:",
    "        p = jwt.decode(token, settings.secret_key, algorithms=[settings.algorithm])",
    "        return p.get('sub')",
    "    except jwt.PyJWTError:",
    "        return None",
])
bullet(doc, "Login mints a token with create_access_token(user.id).")
bullet(doc, "Every protected call sends Authorization: Bearer <token>; the guard decodes it back to a user id.")

# ===== 7 =====
h1("7. The Database Tables (Models)")
body(doc, "Seven tables, each a class in its feature package. On startup "
          "Base.metadata.create_all() issues CREATE TABLE for any missing table.")
make_table(doc, ["Table", "Class", "File"], [
    ["users", "User", "app/auth/models.py"],
    ["books", "Book", "app/books/models.py"],
    ["carts", "Cart", "app/cart/models.py"],
    ["wishlists", "WishList", "app/wishlist/models.py"],
    ["addresses", "Address", "app/address/models.py"],
    ["orders", "Order", "app/orders/models.py"],
    ["feedbacks", "Feedback", "app/feedback/models.py"],
])
body(doc, "Cross-feature links without import cycles: reference the other class by string "
          "name in relationship('Book'), and put the real import behind TYPE_CHECKING. "
          "SQLAlchemy resolves the string later, after every model has been imported.", bold=True)
h2("app/auth/models.py")
code_block(doc, [
    "from __future__ import annotations",
    "from typing import TYPE_CHECKING",
    "from sqlalchemy import Integer, String",
    "from sqlalchemy.orm import Mapped, mapped_column, relationship",
    "from app.core.database import Base, TimestampMixin",
    "if TYPE_CHECKING:",
    "    from app.address.models import Address",
    "    from app.books.models import Book",
    "    from app.cart.models import Cart",
    "    from app.feedback.models import Feedback",
    "    from app.wishlist.models import WishList",
    "",
    "class User(Base, TimestampMixin):",
    "    __tablename__ = 'users'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    role: Mapped[str] = mapped_column(String(255), default='user')  # user|admin",
    "    first_name: Mapped[str] = mapped_column(String(255))",
    "    last_name: Mapped[str] = mapped_column(String(255))",
    "    email: Mapped[str] = mapped_column(String(255), unique=True, index=True)",
    "    phone_no: Mapped[str] = mapped_column(String(255))",
    "    password: Mapped[str] = mapped_column(String(255))   # the HASH",
    "    books/carts/wishlists/addresses/feedbacks:",
    "        Mapped[list[...]] = relationship(back_populates='user')",
])
h2("app/books/models.py")
code_block(doc, [
    "class Book(Base, TimestampMixin):",
    "    __tablename__ = 'books'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    user_id: Mapped[int] = mapped_column(Integer, ForeignKey('users.id'), index=True)",
    "    name: Mapped[str] = mapped_column(String(255))",
    "    description: Mapped[str] = mapped_column(String(1000))",
    "    author: Mapped[str] = mapped_column(String(255))",
    "    image: Mapped[str] = mapped_column(String(255))   # a URL string, not an upload",
    "    price: Mapped[str] = mapped_column('price', String(255))",
    "    quantity: Mapped[int] = mapped_column(Integer, default=0)",
    "    user: Mapped['User'] = relationship(back_populates='books')",
])
h2("carts / wishlists / addresses / orders / feedbacks")
code_block(doc, [
    "class Cart(Base, TimestampMixin):      # carts",
    "    __tablename__ = 'carts'",
    "    id; user_id->users.id; book_id->books.id; book_quantity=1",
    "    user (back_populates='carts'); book",
    "",
    "class WishList(Base, TimestampMixin):  # wishlists",
    "    id; user_id->users.id; book_id->books.id",
    "",
    "class Address(Base, TimestampMixin):   # addresses",
    "    id; user_id->users.id",
    "    address/city/state/landmark: String(255); pincode: Integer",
    "    address_type: String(255) default 'home'",
    "",
    "class Order(Base, TimestampMixin):     # orders",
    "    id; user_id->users.id; book_id->books.id; address_id->addresses.id",
    "    order_id: String(255) nullable",
    "",
    "class Feedback(Base, TimestampMixin):  # feedbacks",
    "    id; user_id->users.id; book_id->books.id",
    "    feedback: String(255); rating: Integer",
])
body(doc, "Each file starts with the same from __future__ / TYPE_CHECKING header, importing "
          "Base, TimestampMixin and the related class names.", italic=True)

# ===== 8 =====
h1("8. Validation Schemas — per feature schemas.py")
body(doc, "A model is a table; a schema is the shape of a request body. A bad body is "
          "rejected with 422 before your code runs.")
h2("app/auth/schemas.py")
code_block(doc, [
    "from pydantic import BaseModel, EmailStr, Field, field_validator",
    "",
    "class RegisterRequest(BaseModel):",
    "    role: str = Field(pattern='^(user|admin)$')",
    "    first_name/last_name: str = Field(min_length=2, max_length=50)",
    "    phone_no: str = Field(min_length=10)",
    "    email: EmailStr = Field(max_length=100)",
    "    password: str = Field(min_length=6)",
    "    confirm_password: str",
    "    @field_validator('confirm_password')",
    "    @classmethod",
    "    def match(cls, v, info):",
    "        if v != info.data.get('password'):",
    "            raise ValueError('confirm_password must match password')",
    "        return v",
    "",
    "class LoginRequest(BaseModel):  email: EmailStr; password: str",
    "# ForgotPasswordRequest, ResetPasswordRequest follow the same idea",
])
h2("app/books/schemas.py")
code_block(doc, [
    "from pydantic import BaseModel, ConfigDict, Field",
    "",
    "class BookCreate(BaseModel):",
    "    name: str = Field(min_length=2, max_length=100)",
    "    description: str = Field(min_length=5, max_length=1000)",
    "    author: str = Field(min_length=5, max_length=300)",
    "    image: str",
    "    price: float = Field(ge=0, alias='Price')",
    "    quantity: int = Field(ge=0)",
    "    model_config = ConfigDict(populate_by_name=True)",
    "",
    "class BookId(BaseModel):    id: int",
    "class AddQuantity(BaseModel): id: int; quantity: int = Field(ge=1)",
    "class SearchRequest(BaseModel): search: str",
])
h2("The smaller schemas")
code_block(doc, [
    "# cart/schemas.py",
    "class BookIdRequest:     book_id: int",
    "class CartIdRequest:     cart_id: int",
    "class WishlistIdRequest: wishlist_id: int",
    "",
    "# address/schemas.py",
    "class AddressCreate(BaseModel):",
    "    address/city/state/landmark: str = Field(min_length=2, ...)",
    "    pincode: int; address_type: str = Field(min_length=2, max_length=100)",
    "class AddressUpdate(AddressCreate): id: int",
    "",
    "# orders/schemas.py",
    "class OrderRequest: name: str; address_id: int; quantity: int = Field(ge=1)",
    "",
    "# feedback/schemas.py",
    "class FeedbackRequest:",
    "    book_id: int; feedback: str = Field(min_length=4); rating: int = Field(ge=1, le=5)",
])

# ===== 9 =====
h1("9. JWT Guards — app/auth/dependencies.py")
body(doc, "Backend-only, so this file has just the JWT guards (no web/session helpers).")
code_block(doc, [
    "from fastapi import Depends, HTTPException, status",
    "from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer",
    "from sqlalchemy.orm import Session",
    "from app.core.database import get_db",
    "from app.core.security import decode_access_token",
    "from .models import User",
    "",
    "bearer_scheme = HTTPBearer(auto_error=False)",
    "",
    "def get_current_api_user(credentials=Depends(bearer_scheme), db=Depends(get_db)):",
    "    if credentials is None:",
    "        raise HTTPException(401, 'Authorization token not found')",
    "    sub = decode_access_token(credentials.credentials)",
    "    if sub is None: raise HTTPException(401, 'Invalid or expired token')",
    "    user = db.get(User, int(sub))",
    "    if user is None: raise HTTPException(401, 'User not found')",
    "    return user",
    "",
    "def require_api_role(role):              # factory -> a role-checking guard",
    "    def checker(user=Depends(get_current_api_user)):",
    "        if user.role != role:",
    "            raise HTTPException(403, f'Only {role}s are allowed to do this')",
    "        return user",
    "    return checker",
])
body(doc, "Using HTTPBearer also makes the Authorize button appear in Swagger (/docs), so you "
          "can paste a token and try protected endpoints right in the browser.", italic=True)

# ===== 10 =====
h1("10. The API Routers — app/<feature>/api.py")
body(doc, "Every endpoint follows the same recipe: schema validates the body -> guard checks "
          "the token/role -> session opens -> model read/write -> commit -> return a dict "
          "(FastAPI turns it into JSON).")
h2("10.1 app/auth/api.py")
code_block(doc, [
    "router = APIRouter(tags=['auth'])",
    "",
    "@router.post('/register', status_code=201)",
    "def register(payload: RegisterRequest, db=Depends(get_db)):",
    "    if db.scalar(select(User).where(User.email == payload.email)):",
    "        raise HTTPException(401, 'The email has already been taken')",
    "    user = User(role=payload.role, first_name=payload.first_name, ...,",
    "                password=hash_password(payload.password))",
    "    db.add(user); db.commit(); db.refresh(user)",
    "    return {'status': 201, 'user_id': user.id}",
    "",
    "@router.post('/login')",
    "def login(payload: LoginRequest, db=Depends(get_db)):",
    "    user = db.scalar(select(User).where(User.email == payload.email))",
    "    if not user or not verify_password(payload.password, user.password):",
    "        raise HTTPException(401, 'Wrong email or password')",
    "    return {'access_token': create_access_token(user.id), 'token_type': 'bearer'}",
    "",
    "# logout (stateless), forgotPassword, resetPassword follow the same shape",
])
h2("10.2 app/books/api.py (admin mutates, anyone logged-in reads)")
code_block(doc, [
    "router = APIRouter(tags=['books'])",
    "",
    "def _serialize(b):",
    "    return {'id': b.id, 'name': b.name, 'author': b.author,",
    "            'image': b.image, 'Price': b.price, 'quantity': b.quantity}",
    "",
    "@router.post('/addingBook', status_code=201)",
    "def adding_book(payload: BookCreate,",
    "                admin=Depends(require_api_role('admin')), db=Depends(get_db)):",
    "    book = Book(user_id=admin.id, name=payload.name, ...,",
    "                price=str(payload.price), quantity=payload.quantity)",
    "    db.add(book); db.commit(); db.refresh(book)      # INSERT into books",
    "    return {'status': 201, 'book': _serialize(book)}",
    "",
    "@router.get('/displayAllBooks')",
    "def display_all(_=Depends(get_current_api_user), db=Depends(get_db)):",
    "    books = db.scalars(select(Book).order_by(Book.name)).all()",
    "    return {'status': 200, 'books': [_serialize(b) for b in books]}",
    "# also: updateBookById, deleteBookById, addQuantityToExistBook,",
    "#       sortPriceLowToHigh, sortPriceHighToLow, searchBookByKeyword",
])
h2("10.3 app/cart/api.py (role 'user'; note the ownership check)")
code_block(doc, [
    "router = APIRouter(tags=['cart'])",
    "user_role = require_api_role('user')",
    "",
    "@router.post('/addBookToCartByBookId', status_code=201)",
    "def add_to_cart(payload: BookIdRequest, user=Depends(user_role), db=Depends(get_db)):",
    "    book = db.get(Book, payload.book_id)",
    "    if not book:           raise HTTPException(404, 'Book not found')",
    "    if book.quantity == 0: raise HTTPException(400, 'OUT OF STOCK')",
    "    if db.scalar(select(Cart).where(Cart.book_id==book.id, Cart.user_id==user.id)):",
    "        raise HTTPException(409, 'Book already added to the cart')",
    "    db.add(Cart(book_id=book.id, user_id=user.id)); db.commit()",
    "    return {'status': 201, 'message': 'Book added to cart Successfully'}",
    "",
    "@router.post('/deleteBookByCartId')",
    "def delete_cart(payload: CartIdRequest, user=Depends(user_role), db=Depends(get_db)):",
    "    cart = db.get(Cart, payload.cart_id)",
    "    if not cart or cart.user_id != user.id:      # never trust the id alone",
    "        raise HTTPException(404, 'Cart not found for this user')",
    "    db.delete(cart); db.commit()",
    "    return {'status': 201, 'message': 'Book deleted from cart Successfully'}",
])
h2("10.4 The rest (same recipe)")
bullet(doc, "wishlist/api.py — addBookToWishlistByBookId, getAllBooksInWishlist, deleteByWishlistId.")
bullet(doc, "address/api.py — addAddress, updateAddress, deleteAddress, getAddress (each write checks address.user_id == user.id).")
bullet(doc, "orders/api.py — placeOrder: find book by name, verify stock, verify address owner, create Order with random order_id, decrement book.quantity.")
bullet(doc, "feedback/api.py — feedback (insert), getAverageRatingByBookId (func.avg(Feedback.rating)).")
body(doc, "Example — placeOrder:")
code_block(doc, [
    "@router.post('/placeOrder', status_code=201)",
    "def place_order(payload: OrderRequest, user=Depends(user_role), db=Depends(get_db)):",
    "    book = db.scalar(select(Book).where(Book.name == payload.name))",
    "    if not book:                         raise HTTPException(404, 'No such book')",
    "    if book.quantity < payload.quantity: raise HTTPException(400, 'Not enough stock')",
    "    address = db.get(Address, payload.address_id)",
    "    if not address or address.user_id != user.id:",
    "        raise HTTPException(404, 'This address id is not available')",
    "    oid = ''.join(random.choices(string.digits + string.ascii_lowercase, k=10))",
    "    db.add(Order(user_id=user.id, book_id=book.id, address_id=address.id, order_id=oid))",
    "    book.quantity -= payload.quantity; db.commit()",
    "    return {'status': 201, 'order_id': oid,",
    "            'total_price': payload.quantity * float(book.price)}",
])

# ===== 11 =====
h1("11. Wire It Together — app/main.py")
body(doc, "Backend-only: no SessionMiddleware, no static mount, no templating. Just the "
          "table-creation hook and the API routers. (FastAPI already returns a clean 422 "
          "JSON for bad bodies, so no custom validation handler is needed.)")
code_block(doc, [
    "from fastapi import FastAPI",
    "from app.core.config import settings",
    "from app.core.database import Base, engine",
    "from app.auth import api as auth_api",
    "from app.books import api as books_api",
    "# ...one import per feature (cart, wishlist, address, orders, feedback)...",
    "",
    "app = FastAPI(title=settings.app_name, version='1.0.0')",
    "",
    "@app.on_event('startup')",
    "def on_startup():",
    "    if settings.create_tables:",
    "        Base.metadata.create_all(bind=engine)   # CREATE TABLE x7",
    "",
    "for module in (auth_api, books_api, cart_api, wishlist_api,",
    "               address_api, orders_api, feedback_api):",
    "    app.include_router(module.router, prefix='/api')",
])
body(doc, "Importing every router here pulls in every models.py, so when create_all runs, "
          "SQLAlchemy knows all 7 tables and the string relationships resolve.", bold=True)

# ===== 12 =====
h1("12. Run + Test")
body(doc, "run.py:")
code_block(doc, [
    "import uvicorn",
    "if __name__ == '__main__':",
    "    uvicorn.run('app.main:app', host='127.0.0.1', port=8000, reload=True)",
])
body(doc, "Create the database once, start MySQL, then run:")
code_block(doc, ["CREATE DATABASE IF NOT EXISTS book_store_product;", "python run.py"])
body(doc, "Open http://127.0.0.1:8000/docs — the interactive Swagger UI lists every endpoint.")
h2("The auth flow with curl")
code_block(doc, [
    "# 1) Register an admin",
    "curl -X POST http://127.0.0.1:8000/api/register -H 'Content-Type: application/json' \\",
    "  -d '{\"role\":\"admin\",\"first_name\":\"Ada\",\"last_name\":\"Lovelace\",",
    "       \"phone_no\":\"9999999999\",\"email\":\"ada@example.com\",",
    "       \"password\":\"secret1\",\"confirm_password\":\"secret1\"}'",
    "",
    "# 2) Log in -> copy the access_token",
    "curl -X POST http://127.0.0.1:8000/api/login -H 'Content-Type: application/json' \\",
    "  -d '{\"email\":\"ada@example.com\",\"password\":\"secret1\"}'",
    "",
    "# 3) Call a protected endpoint with the token",
    "curl http://127.0.0.1:8000/api/displayAllBooks \\",
    "  -H 'Authorization: Bearer <PASTE_TOKEN>'",
])
body(doc, "In Swagger, click Authorize, paste the token once, and every protected endpoint is "
          "unlocked for browser testing.", italic=True)

# ===== 13 =====
h1("13. Endpoint Map")
make_table(doc, ["Method & path", "Role", "Purpose"], [
    ["POST /api/register /login /logout", "-", "Account + JWT"],
    ["POST /api/forgotPassword /resetPassword", "- / token", "Password reset"],
    ["POST /api/addingBook /updateBookById /deleteBookById /addQuantityToExistBook", "admin", "Manage catalogue"],
    ["GET /api/displayAllBooks /sortPriceLowToHigh /sortPriceHighToLow ; POST /searchBookByKeyword", "any", "Browse"],
    ["POST /api/addBookToCartByBookId /deleteBookByCartId /increament /decrement /addBookToCartByWishlistId ; GET /getAllBooksInCart", "user", "Cart"],
    ["POST /api/addBookToWishlistByBookId /deleteBookByWishlistId ; GET /getAllBooksInWishlist", "user", "Wishlist"],
    ["POST /api/addAddress /updateAddress /deleteAddress /getAddress", "user", "Address"],
    ["POST /api/placeOrder", "user", "Orders"],
    ["POST /api/feedback /getAverageRatingByBookId", "user", "Feedback"],
])

# ===== 14 =====
h1("14. Build-Order Checklist")
numbered(doc, "requirements.txt -> install.")
numbered(doc, "app/core/config.py -> .env.")
numbered(doc, "app/core/database.py (engine, Base, TimestampMixin, get_db).")
numbered(doc, "app/core/security.py (hashing + JWT).")
numbered(doc, "Models (tables): auth -> books -> cart -> wishlist -> address -> orders -> feedback.")
numbered(doc, "Schemas per feature.")
numbered(doc, "app/auth/dependencies.py (JWT guards).")
numbered(doc, "API routers per feature (api.py).")
numbered(doc, "app/main.py (table creation + include routers).")
numbered(doc, "run.py -> python run.py -> test in /docs.")

# ===== 15 =====
h1("15. Common Gotchas")
bullet(doc, "401 on every call. Send Authorization: Bearer <token> from POST /api/login. In /docs use Authorize.")
bullet(doc, "403 Forbidden. The endpoint needs a specific role (admin vs user). Register with the right role.")
bullet(doc, "Table missing a new column. create_all only creates missing tables; it never alters existing ones. Drop the table and restart, or use Alembic.")
bullet(doc, "Can't connect to MySQL. Is it running? Do .env credentials match? Does book_store_product exist?")
bullet(doc, "Relationship 'failed to locate a name'. A model never got imported. main.py must import every router so every model loads before first use.")
bullet(doc, "Never commit .env. Git-ignore it; commit .env.example instead.")

out = "Book_Store_FastAPI_API_From_Scratch.docx"
doc.save(out)
print("Saved:", out)

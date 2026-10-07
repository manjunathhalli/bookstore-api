"""Generate a Google-Docs-ready .docx of the *Build From Scratch* guide.

Run:  python build_scratch_doc.py
Output: Book_Store_FastAPI_Build_From_Scratch.docx

Upload the .docx to Google Drive, then right-click -> Open with Google Docs.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK = RGBColor(0x1E, 0x29, 0x3B)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
GREY = RGBColor(0x55, 0x5F, 0x70)
CODE_BG = "F2F3F7"


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, CODE_BG)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Consolas")
    rfonts.set(qn("w:hAnsi"), "Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def body(doc, text, *, italic=False, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(8)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text).font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text).font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)


def make_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].paragraphs[0].add_run(h).bold = True
        for r in hdr[i].paragraphs[0].runs:
            r.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].add_run(val).font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)


def h1(text):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = DARK


def h2(text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = INDIGO


# ---- Title ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Book Store (FastAPI)")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = DARK
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("Build From Scratch — Step by Step")
rs.font.size = Pt(14)
rs.font.color.rgb = INDIGO
intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
ri = intro.add_run("Environment, config, all 7 tables, schemas, guards, every router, "
                   "templates, wiring, and running — in build order.")
ri.italic = True
ri.font.size = Pt(10.5)
ri.font.color.rgb = GREY
doc.add_paragraph()

body(doc, "This is a hands-on build manual. Follow it top to bottom to recreate the whole "
          "project from an empty folder, using the feature-first layout (each domain is its "
          "own package; shared infrastructure lives in app/core; shared auth guards in app/auth). "
          "Type the code yourself — you learn far more than copy-pasting.")

# ===== 0 =====
h1("0. Prerequisites")
make_table(doc, ["Need", "Why", "Check"], [
    ["Python 3.11+", "Modern typing (str | None, Mapped[...])", "python --version"],
    ["MySQL running", "Stores the data (XAMPP is fine)", "XAMPP -> Start MySQL"],
    ["A code editor", "VS Code recommended", "-"],
])

# ===== 1 =====
h1("1. Create the Project Skeleton")
body(doc, "Open PowerShell and run:")
code_block(doc, [
    "mkdir Book_Store_FastAPI",
    "cd Book_Store_FastAPI",
    "",
    "python -m venv .venv",
    ".\\.venv\\Scripts\\Activate.ps1      # prompt shows (.venv)",
    "",
    "mkdir app",
    "mkdir app\\core",
    "mkdir app\\auth app\\books app\\cart app\\wishlist app\\address app\\orders app\\feedback",
    "mkdir app\\dashboard app\\users app\\password",
    "mkdir app\\templates app\\static\\book-covers",
])
body(doc, "Every package needs an __init__.py so Python can import it. Create an empty one "
          "in app and every sub-package:")
code_block(doc, [
    "ni app\\__init__.py",
    "ni app\\core\\__init__.py",
    "foreach ($f in 'auth','books','cart','wishlist','address','orders',",
    "                'feedback','dashboard','users','password')",
    "  { ni app\\$f\\__init__.py }",
])
body(doc, "Target structure (files added over the next steps):")
code_block(doc, [
    "Book_Store_FastAPI/",
    "|-- run.py",
    "|-- requirements.txt",
    "|-- .env  /  .env.example",
    "|-- app/",
    "    |-- main.py",
    "    |-- core/      config.py database.py security.py templating.py",
    "    |-- auth/      models.py schemas.py dependencies.py api.py web.py",
    "    |-- books/     models.py schemas.py api.py web.py",
    "    |-- cart/ wishlist/ address/ orders/ feedback/   (same shape)",
    "    |-- dashboard/ users/ password/   (web.py only)",
    "    |-- templates/ base.html + per-feature .html",
    "    |-- static/book-covers/",
])

# ===== 2 =====
h1("2. Install the Libraries")
body(doc, "Create requirements.txt:")
code_block(doc, [
    "fastapi", "uvicorn[standard]", "sqlalchemy", "pymysql",
    "pydantic[email]", "pydantic-settings", "passlib[bcrypt]", "pyjwt",
    "jinja2", "python-multipart", "itsdangerous",
])
body(doc, "Then install:")
code_block(doc, ["pip install -r requirements.txt"])
make_table(doc, ["Library", "Role"], [
    ["fastapi", "Web framework (routes, validation, dependency injection)."],
    ["uvicorn", "ASGI server that runs the app."],
    ["sqlalchemy", "ORM — Python classes instead of raw SQL."],
    ["pymysql", "Driver SQLAlchemy uses to reach MySQL."],
    ["pydantic / -settings", "Validates request bodies; loads typed settings from .env."],
    ["passlib[bcrypt]", "Hashes passwords (compatible with Laravel $2y$)."],
    ["pyjwt", "Creates/reads the JWT token for the API."],
    ["jinja2", "HTML templating for the web UI."],
    ["python-multipart", "Lets FastAPI read form posts and file uploads."],
    ["itsdangerous", "Signs the session cookie so it can't be tampered with."],
])

# ===== 3 =====
h1("3. Configuration — app/core/config.py")
body(doc, "Everything needs settings (DB credentials, secret key). Read them from .env with "
          "Pydantic so they are typed and validated.")
code_block(doc, [
    "from functools import lru_cache",
    "from pydantic_settings import BaseSettings, SettingsConfigDict",
    "",
    "class Settings(BaseSettings):",
    "    app_name: str = 'Book Store'",
    "    secret_key: str = 'insecure-dev-secret-change-me'",
    "    access_token_expire_minutes: int = 1440",
    "    algorithm: str = 'HS256'",
    "    db_host: str = '127.0.0.1'",
    "    db_port: int = 3306",
    "    db_database: str = 'book_store_product'",
    "    db_username: str = 'root'",
    "    db_password: str = ''",
    "    create_tables: bool = True",
    "    model_config = SettingsConfigDict(env_file='.env', extra='ignore')",
    "",
    "    @property",
    "    def database_url(self) -> str:",
    "        return (f'mysql+pymysql://{self.db_username}:{self.db_password}'",
    "                f'@{self.db_host}:{self.db_port}/{self.db_database}?charset=utf8mb4')",
    "",
    "@lru_cache",
    "def get_settings() -> Settings:",
    "    return Settings()",
    "",
    "settings = get_settings()",
])
body(doc, "Create your secret file .env (and a committable template .env.example):")
code_block(doc, [
    "APP_NAME=Book Store",
    "SECRET_KEY=change-me-to-a-long-random-string",
    "DB_HOST=127.0.0.1",
    "DB_PORT=3306",
    "DB_DATABASE=book_store_product",
    "DB_USERNAME=root",
    "DB_PASSWORD=",
    "CREATE_TABLES=true",
])
body(doc, "Why a database_url property? The rest of the app never assembles the connection "
          "string by hand — it just asks settings.database_url. Field names like DB_PASSWORD "
          "map automatically to db_password.", italic=True)

# ===== 4 =====
h1("4. The Database Engine — app/core/database.py")
body(doc, "Opens the connection pool, gives each request a short-lived session, and defines "
          "Base plus a TimestampMixin every table reuses.")
code_block(doc, [
    "from collections.abc import Generator",
    "from datetime import datetime",
    "from sqlalchemy import DateTime, create_engine, func",
    "from sqlalchemy.orm import (DeclarativeBase, Mapped, Session,",
    "                            mapped_column, sessionmaker)",
    "from .config import settings",
    "",
    "engine = create_engine(settings.database_url, pool_pre_ping=True, future=True)",
    "SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False, future=True)",
    "",
    "class Base(DeclarativeBase):",
    "    pass",
    "",
    "class TimestampMixin:",
    "    created_at: Mapped[datetime | None] = mapped_column(",
    "        DateTime, server_default=func.now())",
    "    updated_at: Mapped[datetime | None] = mapped_column(",
    "        DateTime, server_default=func.now(), onupdate=func.now())",
    "",
    "def get_db() -> Generator[Session, None, None]:",
    "    db = SessionLocal()",
    "    try:",
    "        yield db",
    "    finally:",
    "        db.close()",
])
body(doc, "Key idea — get_db: any endpoint that needs the DB writes db: Session = "
          "Depends(get_db). FastAPI runs get_db, gives you a session, and the finally "
          "guarantees it is closed. You never leak connections.", bold=True)

# ===== 5 =====
h1("5. Passwords + Tokens — app/core/security.py")
code_block(doc, [
    "from datetime import datetime, timedelta, timezone",
    "import jwt",
    "from passlib.context import CryptContext",
    "from .config import settings",
    "",
    "pwd_context = CryptContext(schemes=['bcrypt'], deprecated='auto')",
    "",
    "def hash_password(plain): return pwd_context.hash(plain)",
    "",
    "def verify_password(plain, hashed):",
    "    try:    return pwd_context.verify(plain, hashed)",
    "    except ValueError: return False",
    "",
    "def create_access_token(subject, expires_minutes=None):",
    "    expire = datetime.now(timezone.utc) + timedelta(",
    "        minutes=expires_minutes or settings.access_token_expire_minutes)",
    "    payload = {'sub': str(subject), 'exp': expire}",
    "    return jwt.encode(payload, settings.secret_key, algorithm=settings.algorithm)",
    "",
    "def decode_access_token(token):",
    "    try:",
    "        p = jwt.decode(token, settings.secret_key, algorithms=[settings.algorithm])",
    "        return p.get('sub')",
    "    except jwt.PyJWTError:",
    "        return None",
])
bullet(doc, "Hashing means a database leak never exposes real passwords.")
bullet(doc, "JWT is a signed string proving 'I am user 5' without a server-side session — "
            "the API door uses it. The web door uses a cookie session instead.")

# ===== 6 =====
h1("6. Creating the Database Tables (Models)")
body(doc, "In SQLAlchemy you describe a table as a Python class; on startup "
          "Base.metadata.create_all() issues CREATE TABLE for any table that does not yet "
          "exist. One class = one table, one mapped_column = one column. There are 7 tables, "
          "each owned by its feature package:")
make_table(doc, ["Table", "Class", "Lives in"], [
    ["users", "User", "app/auth/models.py"],
    ["books", "Book", "app/books/models.py"],
    ["carts", "Cart", "app/cart/models.py"],
    ["wishlists", "WishList", "app/wishlist/models.py"],
    ["addresses", "Address", "app/address/models.py"],
    ["orders", "Order", "app/orders/models.py"],
    ["feedbacks", "Feedback", "app/feedback/models.py"],
])
body(doc, "Cross-feature links without circular imports: a User has many Books, but the model "
          "files must not import each other at load time. Reference the other class by string "
          "name in relationship('Book'), and put the real import behind TYPE_CHECKING (used "
          "only by type-checkers). SQLAlchemy resolves the string later, once every model has "
          "been imported.", bold=True)

h2("6.1 users — app/auth/models.py")
code_block(doc, [
    "from __future__ import annotations",
    "from typing import TYPE_CHECKING",
    "from sqlalchemy import Integer, String",
    "from sqlalchemy.orm import Mapped, mapped_column, relationship",
    "from app.core.database import Base, TimestampMixin",
    "",
    "if TYPE_CHECKING:",
    "    from app.address.models import Address",
    "    from app.books.models import Book",
    "    from app.cart.models import Cart",
    "    from app.feedback.models import Feedback",
    "    from app.wishlist.models import WishList",
    "",
    "class User(Base, TimestampMixin):",
    "    __tablename__ = 'users'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    role: Mapped[str] = mapped_column(String(255), default='user')   # user|admin",
    "    first_name: Mapped[str] = mapped_column(String(255))",
    "    last_name: Mapped[str] = mapped_column(String(255))",
    "    email: Mapped[str] = mapped_column(String(255), unique=True, index=True)",
    "    phone_no: Mapped[str] = mapped_column(String(255))",
    "    password: Mapped[str] = mapped_column(String(255))   # stores the HASH",
    "    books:     Mapped[list['Book']]     = relationship(back_populates='user')",
    "    carts:     Mapped[list['Cart']]     = relationship(back_populates='user')",
    "    wishlists: Mapped[list['WishList']] = relationship(back_populates='user')",
    "    addresses: Mapped[list['Address']]  = relationship(back_populates='user')",
    "    feedbacks: Mapped[list['Feedback']] = relationship(back_populates='user')",
])

h2("6.2 books — app/books/models.py")
code_block(doc, [
    "from __future__ import annotations",
    "from typing import TYPE_CHECKING",
    "from sqlalchemy import ForeignKey, Integer, String",
    "from sqlalchemy.orm import Mapped, mapped_column, relationship",
    "from app.core.database import Base, TimestampMixin",
    "if TYPE_CHECKING:",
    "    from app.auth.models import User",
    "",
    "class Book(Base, TimestampMixin):",
    "    __tablename__ = 'books'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    user_id: Mapped[int] = mapped_column(Integer, ForeignKey('users.id'), index=True)",
    "    name: Mapped[str] = mapped_column(String(255))",
    "    description: Mapped[str] = mapped_column(String(1000))",
    "    author: Mapped[str] = mapped_column(String(255))",
    "    image: Mapped[str] = mapped_column(String(255))",
    "    price: Mapped[str] = mapped_column('price', String(255))",
    "    quantity: Mapped[int] = mapped_column(Integer, default=0)",
    "    user: Mapped['User'] = relationship(back_populates='books')",
])
body(doc, "ForeignKey('users.id') makes books.user_id point at users.id. back_populates on "
          "both sides keeps the Python objects in sync.", italic=True)

h2("6.3 carts / wishlists (join tables)")
code_block(doc, [
    "# app/cart/models.py",
    "class Cart(Base, TimestampMixin):",
    "    __tablename__ = 'carts'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    user_id: Mapped[int] = mapped_column(Integer, ForeignKey('users.id'))",
    "    book_id: Mapped[int] = mapped_column(Integer, ForeignKey('books.id'))",
    "    book_quantity: Mapped[int] = mapped_column(Integer, default=1)",
    "    user: Mapped['User'] = relationship(back_populates='carts')",
    "    book: Mapped['Book'] = relationship()",
    "",
    "# app/wishlist/models.py",
    "class WishList(Base, TimestampMixin):",
    "    __tablename__ = 'wishlists'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    user_id: Mapped[int] = mapped_column(Integer, ForeignKey('users.id'))",
    "    book_id: Mapped[int] = mapped_column(Integer, ForeignKey('books.id'))",
    "    user: Mapped['User'] = relationship(back_populates='wishlists')",
    "    book: Mapped['Book'] = relationship()",
])

h2("6.4 addresses / orders / feedbacks")
code_block(doc, [
    "# app/address/models.py",
    "class Address(Base, TimestampMixin):",
    "    __tablename__ = 'addresses'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    user_id: Mapped[int] = mapped_column(Integer, ForeignKey('users.id'))",
    "    address/city/state/landmark: Mapped[str] = mapped_column(String(255))",
    "    pincode: Mapped[int] = mapped_column(Integer)",
    "    address_type: Mapped[str] = mapped_column(String(255), default='home')",
    "    user: Mapped['User'] = relationship(back_populates='addresses')",
    "",
    "# app/orders/models.py",
    "class Order(Base, TimestampMixin):",
    "    __tablename__ = 'orders'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    user_id/book_id/address_id: Mapped[int] = mapped_column(",
    "        Integer, ForeignKey('users.id'/'books.id'/'addresses.id'))",
    "    order_id: Mapped[str | None] = mapped_column(String(255), nullable=True)",
    "    user/book/address: Mapped[...] = relationship()",
    "",
    "# app/feedback/models.py",
    "class Feedback(Base, TimestampMixin):",
    "    __tablename__ = 'feedbacks'",
    "    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)",
    "    user_id: Mapped[int] = mapped_column(Integer, ForeignKey('users.id'))",
    "    book_id: Mapped[int] = mapped_column(Integer, ForeignKey('books.id'))",
    "    feedback: Mapped[str] = mapped_column(String(255))",
    "    rating: Mapped[int] = mapped_column(Integer)",
    "    user: Mapped['User'] = relationship(back_populates='feedbacks')",
    "    book: Mapped['Book'] = relationship()",
])
body(doc, "When do tables get created? Not yet — in step 11, when main.py's startup hook runs "
          "Base.metadata.create_all(bind=engine). By then main.py has imported every router "
          "(which import every model), so SQLAlchemy knows all 7 tables.", italic=True)

# ===== 7 =====
h1("7. Validation Schemas (Pydantic)")
body(doc, "A model is a database table (SQLAlchemy). A schema is the shape of a "
          "request/response (Pydantic). If a body doesn't match, FastAPI returns 422 BEFORE "
          "your code runs. Each feature owns its schemas.py.")
h2("7.1 app/auth/schemas.py")
code_block(doc, [
    "from pydantic import BaseModel, EmailStr, Field, field_validator",
    "",
    "class RegisterRequest(BaseModel):",
    "    role: str = Field(pattern='^(user|admin)$')",
    "    first_name: str = Field(min_length=2, max_length=50)",
    "    last_name: str = Field(min_length=2, max_length=50)",
    "    phone_no: str = Field(min_length=10)",
    "    email: EmailStr = Field(max_length=100)",
    "    password: str = Field(min_length=6)",
    "    confirm_password: str",
    "    @field_validator('confirm_password')",
    "    @classmethod",
    "    def match(cls, v, info):",
    "        if v != info.data.get('password'):",
    "            raise ValueError('confirm_password must match password')",
    "        return v",
    "",
    "class LoginRequest(BaseModel):",
    "    email: EmailStr",
    "    password: str",
    "# ForgotPasswordRequest, ResetPasswordRequest follow the same idea",
])
h2("7.2 app/books/schemas.py")
code_block(doc, [
    "from pydantic import BaseModel, ConfigDict, Field",
    "",
    "class BookCreate(BaseModel):",
    "    name: str = Field(min_length=2, max_length=100)",
    "    description: str = Field(min_length=5, max_length=1000)",
    "    author: str = Field(min_length=5, max_length=300)",
    "    image: str",
    "    price: float = Field(ge=0, alias='Price')",
    "    quantity: int = Field(ge=0)",
    "    model_config = ConfigDict(populate_by_name=True)",
    "",
    "class BookId(BaseModel):     id: int",
    "class AddQuantity(BaseModel): id: int; quantity: int = Field(ge=1)",
    "class SearchRequest(BaseModel): search: str",
    "# BookUpdate adds id + the editable fields",
])
h2("7.3 The smaller schemas (cart / wishlist / address / orders / feedback)")
code_block(doc, [
    "# cart/schemas.py",
    "class BookIdRequest(BaseModel):     book_id: int",
    "class CartIdRequest(BaseModel):     cart_id: int",
    "class WishlistIdRequest(BaseModel): wishlist_id: int",
    "",
    "# address/schemas.py",
    "class AddressCreate(BaseModel):",
    "    address/city/state/landmark: str = Field(min_length=2, ...)",
    "    pincode: int",
    "    address_type: str = Field(min_length=2, max_length=100)",
    "class AddressUpdate(AddressCreate): id: int",
    "",
    "# orders/schemas.py",
    "class OrderRequest(BaseModel): name: str; address_id: int; quantity: int = Field(ge=1)",
    "",
    "# feedback/schemas.py",
    "class FeedbackRequest(BaseModel):",
    "    book_id: int; feedback: str = Field(min_length=4); rating: int = Field(ge=1, le=5)",
])

# ===== 8 =====
h1("8. The Auth Guards — app/auth/dependencies.py")
body(doc, "Dependencies are reusable checks attached to endpoints. They answer 'who are you, "
          "and are you allowed here?' in one place. They live in auth because every other "
          "feature imports them.")
code_block(doc, [
    "from fastapi import Depends, HTTPException, Request, status",
    "from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer",
    "from sqlalchemy.orm import Session",
    "from app.core.database import get_db",
    "from app.core.security import decode_access_token",
    "from .models import User",
    "",
    "class RedirectException(Exception):       # web guards raise this to bounce",
    "    def __init__(self, url, error=None, status=None):",
    "        self.url, self.error, self.status = url, error, status",
    "",
    "bearer_scheme = HTTPBearer(auto_error=False)",
    "",
    "def get_current_api_user(credentials=Depends(bearer_scheme), db=Depends(get_db)):",
    "    if credentials is None:",
    "        raise HTTPException(401, 'Authorization token not found')",
    "    sub = decode_access_token(credentials.credentials)",
    "    if sub is None: raise HTTPException(401, 'Invalid or expired token')",
    "    user = db.get(User, int(sub))",
    "    if user is None: raise HTTPException(401, 'User not found')",
    "    return user",
    "",
    "def require_api_role(role):                # factory -> a role-checking guard",
    "    def checker(user=Depends(get_current_api_user)):",
    "        if user.role != role:",
    "            raise HTTPException(403, f'Only {role}s are allowed to do this')",
    "        return user",
    "    return checker",
    "",
    "def get_optional_web_user(request, db=Depends(get_db)):",
    "    uid = request.session.get('user_id')",
    "    return db.get(User, int(uid)) if uid else None",
    "",
    "def require_web_user(request, db=Depends(get_db)):",
    "    user = get_optional_web_user(request, db)",
    "    if user is None:",
    "        raise RedirectException('/login', error='Please sign in to continue.')",
    "    return user",
    "",
    "def require_web_role(role):",
    "    def checker(user=Depends(require_web_user)):",
    "        if user.role != role:",
    "            raise RedirectException('/dashboard', error='Not allowed there.')",
    "        return user",
    "    return checker",
])
body(doc, "The payoff: an admin-only endpoint just writes "
          "admin: User = Depends(require_api_role('admin')). No auth code is repeated inside "
          "the endpoint — the guard runs first; if it fails the body never executes.", bold=True)

# ===== 9 =====
h1("9. The Routers (the Endpoints)")
body(doc, "Each feature exposes two routers: api.py (JSON, JWT guarded) and web.py "
          "(HTML/redirects, cookie-session guarded). Read the first example slowly.")

h2("9.1 app/auth/api.py — register / login / logout / password")
code_block(doc, [
    "router = APIRouter(tags=['auth'])",
    "",
    "@router.post('/register', status_code=201)",
    "def register(payload: RegisterRequest, db=Depends(get_db)):",
    "    if db.scalar(select(User).where(User.email == payload.email)):",
    "        raise HTTPException(401, 'The email has already been taken')",
    "    user = User(role=payload.role, first_name=payload.first_name, ...,",
    "                password=hash_password(payload.password))",
    "    db.add(user); db.commit(); db.refresh(user)",
    "    return {'status': 201, 'user_id': user.id}",
    "",
    "@router.post('/login')",
    "def login(payload: LoginRequest, db=Depends(get_db)):",
    "    user = db.scalar(select(User).where(User.email == payload.email))",
    "    if not user or not verify_password(payload.password, user.password):",
    "        raise HTTPException(401, 'Wrong email or password')",
    "    return {'access_token': create_access_token(user.id), 'token_type': 'bearer'}",
    "",
    "# logout (stateless), forgotPassword, resetPassword follow the same shape",
])

h2("9.2 app/books/api.py — the canonical CRUD example")
code_block(doc, [
    "router = APIRouter(tags=['books'])",
    "",
    "@router.post('/addingBook', status_code=201)",
    "def adding_book(",
    "    payload: BookCreate,                              # schema validates the body",
    "    admin: User = Depends(require_api_role('admin')), # guard: must be admin",
    "    db: Session = Depends(get_db),                    # DB session, auto-closed",
    "):",
    "    book = Book(user_id=admin.id, name=payload.name, description=payload.description,",
    "                author=payload.author, image=payload.image,",
    "                price=str(payload.price), quantity=payload.quantity)",
    "    db.add(book); db.commit(); db.refresh(book)       # INSERT into `books`",
    "    return {'status': 201, 'book': _serialize(book)}",
    "",
    "@router.get('/displayAllBooks')",
    "def display_all_books(_=Depends(get_current_api_user), db=Depends(get_db)):",
    "    books = db.scalars(select(Book).order_by(Book.name)).all()",
    "    return {'status': 200, 'books': [_serialize(b) for b in books]}",
])
body(doc, "That adding_book function IS the whole architecture in miniature: schema -> guard "
          "-> session -> model -> DB. Every other endpoint is a variation. Books also has "
          "updateBookById, deleteBookById, addQuantityToExistBook, sortPriceLowToHigh, "
          "sortPriceHighToLow, searchBookByKeyword.", italic=True)

h2("9.3 The remaining API routers (same recipe)")
bullet(doc, "cart/api.py — addBookToCartByBookId, getAllBooksInCart, deleteBookByCartId, "
            "increament/decrement quantity, addBookToCartByWishlistId. Guard: role 'user'.")
bullet(doc, "wishlist/api.py — addBookToWishlistByBookId, getAllBooksInWishlist, deleteByWishlistId.")
bullet(doc, "address/api.py — addAddress, updateAddress, deleteAddress, getAddress.")
bullet(doc, "orders/api.py — placeOrder (stock check, decrement books.quantity, random order_id).")
bullet(doc, "feedback/api.py — feedback, getAverageRatingByBookId (func.avg(Feedback.rating)).")
body(doc, "The 'owns this row' check used throughout the user features:")
code_block(doc, [
    "@router.post('/deleteBookByCartId')",
    "def delete_book_by_cart_id(payload: CartIdRequest,",
    "                           user=Depends(require_api_role('user')), db=Depends(get_db)):",
    "    cart = db.get(Cart, payload.cart_id)",
    "    if not cart or cart.user_id != user.id:        # never trust the id alone",
    "        raise HTTPException(404, 'Cart not found for this user')",
    "    db.delete(cart); db.commit()",
    "    return {'status': 201, 'message': 'Book deleted from cart Successfully'}",
])

h2("9.4 The web routers — app/<feature>/web.py")
body(doc, "The web version returns a redirect + flash message instead of JSON, and reads form "
          "fields (Form(...)) instead of a JSON schema. Example — books/web.py store with a "
          "cover upload:")
code_block(doc, [
    "@router.post('')",
    "def store(request: Request,",
    "          name: str = Form(...), author: str = Form(...), description: str = Form(...),",
    "          Price: float = Form(...), quantity: int = Form(...),",
    "          image: UploadFile = File(...),",
    "          user: User = Depends(require_web_role('admin')), db=Depends(get_db)):",
    "    image_url = _save_image(image)                       # write under static/",
    "    db.add(Book(user_id=user.id, name=name, author=author, description=description,",
    "                image=image_url, price=str(Price), quantity=quantity))",
    "    db.commit()",
    "    flash(request, f'Book \"{name}\" created successfully.')  # one-time message",
    "    return RedirectResponse('/books', status_code=303)       # Post/Redirect/Get",
])
body(doc, "Why redirect after a POST (PRG)? So a browser refresh does not resubmit the form. "
          "The success message survives via a one-time 'flash' in the signed session cookie.", bold=True)
make_table(doc, ["File", "Prefix", "Pages / actions"], [
    ["auth/web.py", "(root)", "/login /register /logout, / -> /dashboard"],
    ["dashboard/web.py", "(root)", "/dashboard (admin vs user counts)"],
    ["books/web.py", "/books", "list, create, edit, update, add-quantity, delete"],
    ["cart/web.py", "/cart", "list, add, increment, decrement, delete"],
    ["wishlist/web.py", "/wishlist", "list, add, move-to-cart, delete"],
    ["address/web.py", "/address", "list, create, edit, update, delete"],
    ["orders/web.py", "/orders", "list (own/all), place order"],
    ["feedback/web.py", "/feedback", "submit feedback, show average"],
    ["users/web.py", "/users", "admin-only directory of accounts"],
    ["password/web.py", "/password", "forgot (token) + reset"],
])

# ===== 10 =====
h1("10. Web Rendering — app/core/templating.py + templates/")
body(doc, "templating.py sets up Jinja2 and a render() helper that always injects the current "
          "user + any flash messages.")
code_block(doc, [
    "TEMPLATES_DIR = Path(__file__).resolve().parents[1] / 'templates'",
    "templates = Jinja2Templates(directory=str(TEMPLATES_DIR))",
    "",
    "def flash(request, message, category='status'):",
    "    request.session.setdefault('_flash', {})[category] = message",
    "",
    "def render(request, template, context=None, *, user=None, status_code=200):",
    "    fl = request.session.pop('_flash', {})",
    "    base = {'app_name': settings.app_name, 'current_user': user,",
    "            'status': fl.get('status'), 'error': fl.get('error'),",
    "            'errors': request.session.pop('_errors', []),",
    "            'old': request.session.pop('_old', {})}",
    "    if context: base.update(context)",
    "    return templates.TemplateResponse(request, template, base, status_code=status_code)",
])
body(doc, "base.html is the shared skeleton; every page extends it:")
code_block(doc, [
    "<!doctype html><html><head><title>{{ app_name }}</title></head><body>",
    "  <nav>",
    "    {% if current_user %} Hi {{ current_user.first_name }}",
    "      <a href='/dashboard'>Dashboard</a> <a href='/books'>Books</a>",
    "    {% else %} <a href='/login'>Login</a> {% endif %}",
    "  </nav>",
    "  {% if status %}<p class='ok'>{{ status }}</p>{% endif %}",
    "  {% if error %}<p class='err'>{{ error }}</p>{% endif %}",
    "  {% block content %}{% endblock %}",
    "</body></html>",
])
body(doc, "A page just fills the content block, e.g. books/index.html:")
code_block(doc, [
    "{% extends 'base.html' %}",
    "{% block content %}",
    "  {% for book in books %}",
    "    <div class='card'>{{ book.name }} - Rs {{ book.price }}</div>",
    "  {% endfor %}",
    "{% endblock %}",
])

# ===== 11 =====
h1("11. Wire It Together — app/main.py")
body(doc, "Built last because it imports everything. It creates the app, enables the session "
          "cookie, mounts static files, CREATES THE TABLES on startup, registers the "
          "web-redirect/validation exception handlers, and includes every router.")
code_block(doc, [
    "app = FastAPI(title=f'{settings.app_name} API', version='1.0.0')",
    "app.add_middleware(SessionMiddleware, secret_key=settings.secret_key,",
    "                   max_age=60*60*24*7)",
    "app.mount('/static', StaticFiles(directory=str(STATIC_DIR)), name='static')",
    "",
    "@app.on_event('startup')",
    "def on_startup():",
    "    if settings.create_tables:",
    "        Base.metadata.create_all(bind=engine)   # CREATE TABLE x7",
    "",
    "@app.exception_handler(RedirectException)",
    "async def _redirect(request, exc):",
    "    if exc.error:  flash(request, exc.error, 'error')",
    "    return RedirectResponse(exc.url, status_code=303)",
    "",
    "@app.exception_handler(RequestValidationError)",
    "async def _validation(request, exc):",
    "    if request.url.path.startswith('/api'):",
    "        return JSONResponse(status_code=422, content={'detail': exc.errors()})",
    "    flash_errors(request, [e['msg'] for e in exc.errors()])",
    "    return RedirectResponse(request.headers.get('referer', '/dashboard'), 303)",
    "",
    "from app.auth import api as auth_api, web as auth_web   # ...one per feature...",
    "",
    "for m in (auth_api, books_api, cart_api, wishlist_api,",
    "          address_api, orders_api, feedback_api):",
    "    app.include_router(m.router, prefix='/api')          # API door",
    "for m in (auth_web, dashboard_web, books_web, cart_web, wishlist_web,",
    "          address_web, orders_web, feedback_web, users_web, password_web):",
    "    app.include_router(m.router)                         # web door",
])
body(doc, "This import block is also what makes table creation work: importing every api/web "
          "module pulls in every models.py, so by the time create_all runs, SQLAlchemy's "
          "metadata holds all 7 tables and their relationships resolve.", bold=True)

# ===== 12 =====
h1("12. Run It")
body(doc, "Create run.py in the project root:")
code_block(doc, [
    "import uvicorn",
    "if __name__ == '__main__':",
    "    uvicorn.run('app.main:app', host='127.0.0.1', port=8000, reload=True)",
])
body(doc, "Make sure the database exists once:")
code_block(doc, ["CREATE DATABASE IF NOT EXISTS book_store_product;"])
body(doc, "Start MySQL (XAMPP), then run:")
code_block(doc, ["python run.py"])
body(doc, "Open the Web UI at http://127.0.0.1:8000/ and Swagger at "
          "http://127.0.0.1:8000/docs. The first run creates all 7 tables automatically. "
          "Confirm them:")
code_block(doc, [
    "USE book_store_product;",
    "SHOW TABLES;",
    "-- users, books, carts, wishlists, addresses, orders, feedbacks",
])

# ===== 13 =====
h1("13. Follow One Full Request (the 'aha' moment)")
body(doc, "A logged-in admin adds a book through the API:", bold=True)
numbered(doc, "Client -> POST /api/addingBook with a JSON body + Authorization: Bearer <token>.")
numbered(doc, "main.py routed /api/* to the books API router -> matches adding_book.")
numbered(doc, "FastAPI sees payload: BookCreate -> validates the JSON. Bad data => auto-422.")
numbered(doc, "Depends(require_api_role('admin')) -> guard: decode token, load user, check role. Not admin => 403.")
numbered(doc, "Depends(get_db) -> opens a DB session.")
numbered(doc, "Your code builds a Book model, db.add + db.commit => SQLAlchemy emits INSERT via PyMySQL => MySQL stores the row in books.")
numbered(doc, "You return a dict => FastAPI serializes to JSON => client gets 201 Created.")
numbered(doc, "The session is closed automatically (get_db's finally).")
body(doc, "Every layer you built appears exactly once. That is the whole app.", italic=True)

# ===== 14 =====
h1("14. Build-Order Checklist")
numbered(doc, "requirements.txt -> install.")
numbered(doc, "app/core/config.py -> .env.")
numbered(doc, "app/core/database.py (engine, Base, TimestampMixin, get_db).")
numbered(doc, "app/core/security.py (hashing + JWT).")
numbered(doc, "Models (tables): auth -> books -> cart -> wishlist -> address -> orders -> feedback.")
numbered(doc, "Schemas per feature.")
numbered(doc, "app/auth/dependencies.py (guards).")
numbered(doc, "API routers per feature (api.py).")
numbered(doc, "Web routers per feature (web.py).")
numbered(doc, "app/core/templating.py + app/templates/.")
numbered(doc, "app/main.py (wiring + table creation on startup).")
numbered(doc, "run.py -> python run.py.")

# ===== 15 =====
h1("15. Common Gotchas")
bullet(doc, "'Table is missing my new column.' create_all only CREATES missing tables; it does NOT alter existing ones. Drop the table and restart while learning. Real apps use migrations (Alembic).")
bullet(doc, "Can't connect to MySQL. Is MySQL running? Do .env credentials match? Does book_store_product exist?")
bullet(doc, "401 on API calls. Log in (POST /api/login), copy access_token, send it as Authorization: Bearer <token>. In /docs click Authorize.")
bullet(doc, "Web pages bounce to /login. The web door uses the cookie session, not the JWT. Log in through the web form.")
bullet(doc, "RuntimeError: Form data requires 'python-multipart'. Install it (it's in requirements.txt).")
bullet(doc, "Relationship error 'failed to locate a name Book'. A model was never imported. Ensure main.py imports every router so every model loads before first use.")
bullet(doc, "Never commit .env. Add it to .gitignore; commit .env.example instead.")

out = "Book_Store_FastAPI_Build_From_Scratch.docx"
doc.save(out)
print("Saved:", out)
